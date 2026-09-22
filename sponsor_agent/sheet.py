"""The Excel workbook is the single source of truth.

Every entrypoint opens it, mutates it and saves it. Reopening per write is slow
but means no entrypoint can ever hold a stale in-memory copy.

# ponytail: open/save per write, O(rows) scans. Fine to a few thousand rows;
# move to SQLite if the sheet ever gets big enough to notice.
"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook

from sponsor_agent import config

PROSPECTS = "Prospects"
LOG = "Log"

COLUMNS = [
    "company",
    "contact_name",
    "contact_email",
    "website",
    "notes",
    "research_notes",
    "status",
    "touch_number",
    "draft_subject",
    "draft_body",
    "thread_id",
    "sent_at",
    "replied_at",
    "reply_snippet",
    "last_updated",
    "last_error",
]
LOG_COLUMNS = ["timestamp", "node", "prospect", "outcome"]

NEW, DRAFT, NEEDS_REVIEW, SENT, REPLIED = (
    "NEW",
    "DRAFT",
    "NEEDS_REVIEW",
    "SENT",
    "REPLIED",
)


def now() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def path() -> Path:
    return config.WORKBOOK


def ensure_workbook() -> Path:
    p = path()
    if p.exists():
        return p
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = PROSPECTS
    ws.append(COLUMNS)
    wb.create_sheet(LOG).append(LOG_COLUMNS)
    wb.save(p)
    return p


def _open():
    ensure_workbook()
    return load_workbook(path())


def read_rows() -> list[dict[str, Any]]:
    """Every Prospects row as a dict, plus `_row` (1-based worksheet index)."""
    wb = _open()
    ws = wb[PROSPECTS]
    rows = []
    for i, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v not in (None, "") for v in values):
            continue
        row = {c: ("" if v is None else v) for c, v in zip(COLUMNS, values)}
        row["_row"] = i
        rows.append(row)
    wb.close()
    return rows


def append(record: dict[str, Any]) -> int:
    wb = _open()
    ws = wb[PROSPECTS]
    record = {**record, "last_updated": now()}
    ws.append([record.get(c, "") for c in COLUMNS])
    wb.save(path())
    idx = ws.max_row
    wb.close()
    return idx


def update(row_idx: int, **fields: Any) -> None:
    wb = _open()
    ws = wb[PROSPECTS]
    fields.setdefault("last_updated", now())
    for key, value in fields.items():
        if key not in COLUMNS:
            raise KeyError(f"unknown column: {key}")
        ws.cell(row=row_idx, column=COLUMNS.index(key) + 1, value=value)
    wb.save(path())
    wb.close()


def log(node: str, prospect: str, outcome: str) -> None:
    wb = _open()
    wb[LOG].append([now(), node, prospect, outcome])
    wb.save(path())
    wb.close()


def log_rows() -> list[dict[str, Any]]:
    wb = _open()
    rows = [
        dict(zip(LOG_COLUMNS, values))
        for values in wb[LOG].iter_rows(min_row=2, values_only=True)
        if any(v is not None for v in values)
    ]
    wb.close()
    return rows


# --- queries the guardrails depend on -------------------------------------


def _norm(email: Any) -> str:
    return str(email or "").strip().lower()


def already_sent(contact_email: str, touch_number: Any, rows: Iterable[dict] | None = None) -> bool:
    """Dedup key: this exact contact has already had this touch delivered."""
    rows = read_rows() if rows is None else rows
    target = _norm(contact_email)
    for r in rows:
        if _norm(r["contact_email"]) != target:
            continue
        if str(r["touch_number"]) != str(touch_number):
            continue
        if r["status"] in (SENT, REPLIED):
            return True
    return False


def has_replied(contact_email: str, rows: Iterable[dict] | None = None) -> bool:
    """A reply anywhere retires this contact from every future automated touch."""
    rows = read_rows() if rows is None else rows
    target = _norm(contact_email)
    return any(_norm(r["contact_email"]) == target and r["status"] == REPLIED for r in rows)


def known_emails(rows: Iterable[dict] | None = None) -> set[str]:
    rows = read_rows() if rows is None else rows
    return {_norm(r["contact_email"]) for r in rows if r["contact_email"]}
