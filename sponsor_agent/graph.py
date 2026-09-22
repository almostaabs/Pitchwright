"""The LangGraph pipeline.

    ingest -> research -> next_prospect -> personalize -> validate -+-> send ---------+
                               ^                           ^        |                |
                               |                           +--------+ (retry)        |
                               |                                    |                |
                               +------------------------------------+ needs_review --+

Only two outcomes ever wait on a human: NEEDS_REVIEW and REPLIED. Everything
else - drafting, validating, sending, follow-ups - happens unattended.
"""
from __future__ import annotations

import csv
from typing import Any, Optional, TypedDict

from langgraph.graph import END, StateGraph

from sponsor_agent import config, gmail_client, llm, sheet
from sponsor_agent import research as research_mod

INPUT_FIELDS = {
    "company": "company",
    "organisation": "company",
    "organization": "company",
    "contact_name": "contact_name",
    "name": "contact_name",
    "contact": "contact_name",
    "contact_email": "contact_email",
    "email": "contact_email",
    "website": "website",
    "url": "website",
    "site": "website",
    "notes": "notes",
}


class State(TypedDict, total=False):
    queue: list[dict]
    current: Optional[dict]
    attempts: int
    feedback: list[str]
    verdict: str
    sent: int
    facts: dict


# --- input file ------------------------------------------------------------


def _normalize(record: dict) -> dict:
    out: dict[str, str] = {}
    for key, value in record.items():
        canonical = INPUT_FIELDS.get(str(key or "").strip().lower())
        if canonical and value not in (None, ""):
            out[canonical] = str(value).strip()
    return out


def read_input(path=None) -> list[dict]:
    """Rows from the operator's own prospect list. CSV or Excel."""
    path = path or config.INPUT_FILE
    if not path.exists():
        return []

    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
        wb.close()
        if not rows:
            return []
        headers = [str(h or "") for h in rows[0]]
        return [_normalize(dict(zip(headers, values))) for values in rows[1:]]

    with path.open(newline="", encoding="utf-8-sig") as handle:
        return [_normalize(row) for row in csv.DictReader(handle)]


# --- nodes -----------------------------------------------------------------


def ingest(state: State) -> dict:
    records = read_input()
    if not records:
        sheet.log("ingest", "-", f"no new prospects found in {config.INPUT_FILE}")
        return {}

    known = sheet.known_emails()
    added = 0
    for record in records:
        company = record.get("company", "(unnamed)")
        email = record.get("contact_email", "")
        if "@" not in email:
            sheet.log("ingest", company, "skipped: missing or malformed contact_email")
            continue
        if email.lower() in known:
            continue
        sheet.append(
            {
                "company": company,
                "contact_name": record.get("contact_name", ""),
                "contact_email": email,
                "website": record.get("website", ""),
                "notes": record.get("notes", ""),
                "status": sheet.NEW,
                "touch_number": 1,
            }
        )
        known.add(email.lower())
        added += 1
        sheet.log("ingest", company, "added as NEW")

    sheet.log("ingest", "-", f"{added} new prospect(s) from {len(records)} input row(s)")
    return {}


def research(state: State) -> dict:
    """Gather facts for NEW rows, then build the run queue."""
    for row in sheet.read_rows():
        if row["status"] != sheet.NEW or str(row["research_notes"]).strip():
            continue
        notes = research_mod.research(row)
        sheet.update(row["_row"], research_notes=notes)
        outcome = (
            "no data found - will pitch from the fact sheet alone"
            if notes == research_mod.NO_DATA
            else "facts gathered"
        )
        sheet.log("research", row["company"], outcome)

    rows = sheet.read_rows()
    queue = [
        row
        for row in rows
        if row["status"] in (sheet.NEW, sheet.DRAFT)
        and not sheet.has_replied(row["contact_email"], rows)
    ]
    sheet.log("research", "-", f"{len(queue)} prospect(s) queued for drafting")
    return {"queue": queue}


def next_prospect(state: State) -> dict:
    queue = list(state.get("queue") or [])
    sent = state.get("sent", 0)

    if sent >= config.SEND_CAP:
        if queue:
            sheet.log(
                "send_cap",
                "-",
                f"stopped: cap of {config.SEND_CAP} sends reached, "
                f"{len(queue)} prospect(s) left for the next run",
            )
        return {"current": None, "queue": []}

    current = queue.pop(0) if queue else None
    return {"current": current, "queue": queue, "attempts": 0, "feedback": []}


def personalize(state: State) -> dict:
    prospect = state["current"]
    attempt = state.get("attempts", 0) + 1
    try:
        draft = llm.draft_email(
            prospect,
            state["facts"],
            prospect.get("research_notes", ""),
            feedback=state.get("feedback") or None,
            previous=prospect.get("_previous"),
        )
        prospect["draft_subject"] = draft.subject
        prospect["draft_body"] = draft.body
        prospect.pop("_draft_error", None)
        sheet.update(
            prospect["_row"],
            draft_subject=draft.subject,
            draft_body=draft.body,
            status=sheet.DRAFT,
            last_error="",
        )
        sheet.log("personalize", prospect["company"], f"draft written (attempt {attempt})")
    except Exception as exc:  # API error, bad key, malformed output
        # Blank the draft so a stale one from an earlier attempt can never be sent.
        prospect["draft_subject"] = ""
        prospect["draft_body"] = ""
        prospect["_draft_error"] = f"{type(exc).__name__}: {exc}"
        sheet.log("personalize", prospect["company"], f"FAILED (attempt {attempt}): {exc}")
    return {}


def validate(state: State) -> dict:
    prospect = state["current"]

    if prospect.get("_draft_error"):
        verdict = llm.Verdict(
            passed=False, errors=["drafting failed: " + prospect["_draft_error"]]
        )
    else:
        try:
            verdict = llm.validate_draft(
                prospect.get("draft_subject", ""),
                prospect.get("draft_body", ""),
                prospect,
                state["facts"],
                prospect.get("research_notes", ""),
            )
        except Exception as exc:
            verdict = llm.Verdict(
                passed=False, errors=[f"validator error: {type(exc).__name__}: {exc}"]
            )

    if verdict.passed:
        sheet.log("validate", prospect["company"], "passed")
        return {"verdict": "send"}

    attempts = state.get("attempts", 0) + 1
    problems = list(verdict.errors) + [
        "Unsupported claim, not traceable to the fact sheet or research: " + claim
        for claim in verdict.unsupported_claims
    ]
    problems = problems or ["validation failed with no reason given"]
    sheet.log(
        "validate",
        prospect["company"],
        f"failed (attempt {attempts}): " + "; ".join(problems)[:300],
    )

    over_budget = attempts > config.MAX_VALIDATION_RETRIES
    return {
        "attempts": attempts,
        "feedback": problems,
        "verdict": "needs_review" if over_budget else "personalize",
    }


def send(state: State) -> dict:
    """Passing validate means sending. No approval gate - only the guardrails."""
    prospect = state["current"]
    company = prospect["company"]
    email = prospect["contact_email"]
    touch = prospect.get("touch_number", 1)
    rows = sheet.read_rows()

    # A reply that landed mid-run retires this contact immediately.
    if sheet.has_replied(email, rows):
        sheet.update(
            prospect["_row"],
            status=sheet.NEEDS_REVIEW,
            last_error="not sent: this contact replied, the conversation is yours now",
        )
        sheet.log("send", company, "skipped: contact has replied")
        return {}

    if sheet.already_sent(email, touch, rows):
        sheet.update(
            prospect["_row"],
            status=sheet.NEEDS_REVIEW,
            last_error=f"not sent: touch {touch} to {email} was already delivered (duplicate row)",
        )
        sheet.log("send", company, f"skipped: duplicate of touch {touch}")
        return {}

    if config.DRY_RUN:
        sheet.log("send", company, f"DRY_RUN: would have sent touch {touch} to {email}")
        return {}

    try:
        result = gmail_client.send_email(
            email,
            prospect["draft_subject"],
            prospect["draft_body"],
            thread_id=prospect.get("_thread_id") or None,
        )
    except Exception as exc:
        sheet.update(
            prospect["_row"],
            status=sheet.NEEDS_REVIEW,
            last_error=f"send failed: {type(exc).__name__}: {exc}"[:500],
        )
        sheet.log("send", company, f"FAILED: {exc}")
        return {}

    sheet.update(
        prospect["_row"],
        status=sheet.SENT,
        sent_at=sheet.now(),
        thread_id=result["threadId"],
        last_error="",
    )
    sheet.log("send", company, f"sent touch {touch}, thread {result['threadId']}")
    return {"sent": state.get("sent", 0) + 1}


def needs_review(state: State) -> dict:
    prospect = state["current"]
    reason = "; ".join(state.get("feedback") or ["unspecified"])
    message = f"failed validation after {config.MAX_VALIDATION_RETRIES} retries: {reason}"
    sheet.update(prospect["_row"], status=sheet.NEEDS_REVIEW, last_error=message[:500])
    sheet.log("needs_review", prospect["company"], message[:200])
    return {}


# --- wiring ----------------------------------------------------------------


def build(entry: str = "ingest"):
    """entry="ingest" for the main pipeline, "next_prospect" for a prebuilt queue."""
    graph = StateGraph(State)

    if entry == "ingest":
        graph.add_node("ingest", ingest)
        graph.add_node("research", research)
        graph.add_edge("ingest", "research")
        graph.add_edge("research", "next_prospect")

    graph.add_node("next_prospect", next_prospect)
    graph.add_node("personalize", personalize)
    graph.add_node("validate", validate)
    graph.add_node("send", send)
    graph.add_node("needs_review", needs_review)

    graph.add_conditional_edges(
        "next_prospect",
        lambda s: "personalize" if s.get("current") else "done",
        {"personalize": "personalize", "done": END},
    )
    graph.add_edge("personalize", "validate")
    graph.add_conditional_edges(
        "validate",
        lambda s: s.get("verdict", "needs_review"),
        {"send": "send", "personalize": "personalize", "needs_review": "needs_review"},
    )
    graph.add_edge("send", "next_prospect")
    graph.add_edge("needs_review", "next_prospect")

    graph.set_entry_point(entry)
    return graph.compile()


def run(
    entry: str = "ingest", queue: list[dict] | None = None, facts: dict | None = None
) -> dict:
    sheet.ensure_workbook()
    facts = facts if facts is not None else llm.load_fact_sheet()
    state: State = {
        "queue": queue or [],
        "current": None,
        "attempts": 0,
        "feedback": [],
        "sent": 0,
        "facts": facts,
    }
    # Each prospect costs at most ~10 supersteps and the queue only ever shrinks,
    # so this bound is a safety valve, not a real limit on the run.
    budget = len(queue or []) + len(sheet.read_rows()) + len(read_input()) + 10
    return build(entry).invoke(state, config={"recursion_limit": 10 * budget})


# --- reporting -------------------------------------------------------------


def summary() -> dict[str, int]:
    counts: dict[str, int] = {}
    for row in sheet.read_rows():
        counts[row["status"]] = counts.get(row["status"], 0) + 1
    return counts


def print_summary(title: str, sent: Any) -> None:
    counts = summary()
    print(f"\n{title}: {sent} email(s) sent this run.")
    print("Workbook: " + str(config.WORKBOOK))
    for status in (sheet.NEW, sheet.DRAFT, sheet.SENT, sheet.NEEDS_REVIEW, sheet.REPLIED):
        print(f"  {status:<13} {counts.get(status, 0)}")
    waiting = counts.get(sheet.NEEDS_REVIEW, 0) + counts.get(sheet.REPLIED, 0)
    if waiting:
        print(
            f"\n{waiting} row(s) need you: NEEDS_REVIEW (read last_error) "
            "and REPLIED (a real conversation - reply yourself)."
        )
    else:
        print("\nNothing is waiting on you.")
